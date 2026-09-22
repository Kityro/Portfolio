from fastapi import FastAPI, Depends, HTTPException, UploadFile, File, Response
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from fastapi.responses import FileResponse
from sqlalchemy.orm import Session
from typing import List
import models, schemas, database, services
import socket
import uvicorn

# Initialize Database
models.Base.metadata.create_all(bind=database.engine)

app = FastAPI(title="Portfólio Otávio", description="Backend para Consulta de CPF e Crédito")

# Configure CORS
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

@app.get("/")
def read_root():
    return FileResponse("index.html")

@app.post("/consult", response_model=schemas.ConsultationResponse)
def create_consultation(request: schemas.CPFRequest, db: Session = Depends(database.get_db)):
    analysis_data = services.perform_credit_analysis(request.cpf)
    model_columns = models.Consultation.__table__.columns.keys()
    filtered_data = {k: v for k, v in analysis_data.items() if k in model_columns}
    db_consultation = models.Consultation(**filtered_data)
    db.add(db_consultation)
    db.commit()
    db.refresh(db_consultation)
    return db_consultation

@app.get("/history", response_model=List[schemas.ConsultationResponse])
def get_history(db: Session = Depends(database.get_db)):
    return db.query(models.Consultation).order_by(models.Consultation.created_at.desc()).all()

@app.patch("/consult/{consult_id}", response_model=schemas.ConsultationResponse)
def update_consultation(consult_id: int, update_data: schemas.ConsultationUpdate, db: Session = Depends(database.get_db)):
    db_item = db.query(models.Consultation).filter(models.Consultation.id == consult_id).first()
    if not db_item:
        raise HTTPException(status_code=404, detail="Consultation not found")
    
    for key, value in update_data.dict(exclude_unset=True).items():
        setattr(db_item, key, value)
    
    db.commit()
    db.refresh(db_item)
    return db_item

@app.delete("/consult/{consult_id}")
def delete_consultation(consult_id: int, db: Session = Depends(database.get_db)):
    db_item = db.query(models.Consultation).filter(models.Consultation.id == consult_id).first()
    if not db_item:
        raise HTTPException(status_code=404, detail="Consultation not found")
    
    db.delete(db_item)
    db.commit()
    return {"message": "Deleted successfully"}

@app.post("/consult/batch")
async def batch_consultation(file: UploadFile = File(...), db: Session = Depends(database.get_db)):
    import io
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from datetime import datetime, timedelta

    content = await file.read()
    text = content.decode("utf-8", errors="replace")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resultado Lote"

    # --- Estilos ---
    header_font   = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
    header_fill   = PatternFill("solid", fgColor="1A472A")   # verde escuro
    row_fill_alt  = PatternFill("solid", fgColor="F0FFF4")   # verde bem claro (linhas pares)
    border_side   = Side(style="thin", color="CCCCCC")
    cell_border   = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)

    center = Alignment(horizontal="center", vertical="center")
    left   = Alignment(horizontal="left",   vertical="center")
    right  = Alignment(horizontal="right",  vertical="center")

    # --- Cabeçalho ---
    headers = ["Nº", "CPF", "Nome", "Data de Nascimento",
               "CPF Válido?", "Restrição no Nome?", "Valor da Dívida (R$)", "Classificação da Dívida", "Local da Dívida"]
    ws.append(headers)

    for col_idx, _ in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font      = header_font
        cell.fill      = header_fill
        cell.alignment = center
        cell.border    = cell_border

    # --- Larguras fixas por coluna (em caracteres) ---
    col_widths = [6, 18, 38, 22, 14, 22, 22, 24, 24]
    for i, w in enumerate(col_widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    ws.row_dimensions[1].height = 22

    # --- Dados ---
    row_num = 1
    for line in text.splitlines():
        cpf_clean = "".join(filter(str.isdigit, line.strip()))
        if not cpf_clean:
            continue

        is_valid = services.is_valid_cpf(cpf_clean)

        if len(cpf_clean) == 11:
            cpf_fmt = f"{cpf_clean[:3]}.{cpf_clean[3:6]}.{cpf_clean[6:9]}-{cpf_clean[9:]}"
        else:
            cpf_fmt = cpf_clean

        if not is_valid:
            row_data = [row_num, cpf_fmt, "—", "—", "NÃO", "—", "—", "CPF INVÁLIDO", "—"]
        else:
            a = services.perform_credit_analysis(cpf_clean)
            # Store in database to record each request
            model_columns = models.Consultation.__table__.columns.keys()
            filtered_data = {k: v for k, v in a.items() if k in model_columns}
            db_consultation = models.Consultation(**filtered_data)
            db.add(db_consultation)
            db.commit()
            db.refresh(db_consultation)

            tem_restricao   = a.get("restriction") == "RESTRIÇÃO ATIVA"
            restricao_label = "SIM" if tem_restricao else "NÃO"
            debt_amount     = a.get("debt_amount", 0)
            local_divida    = a.get("debt_location", "—")

            if tem_restricao:
                debt_fmt      = f"R$ {debt_amount:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
                classificacao = a.get("debt_class") or ("ACIMA DE MIL" if debt_amount > 1000 else "ABAIXO DE MIL")
            else:
                debt_fmt      = "R$ 0,00"
                classificacao = "SEM RESTRIÇÃO"
                local_divida  = "NADA CONSTA"

            row_data = [row_num, cpf_fmt, a.get("name", "—"), a.get("birth_date", "—"),
                        "SIM", restricao_label, debt_fmt, classificacao, local_divida]

        ws.append(row_data)
        excel_row = row_num + 1  # +1 porque linha 1 é cabeçalho

        # Linha alternada
        fill = row_fill_alt if row_num % 2 == 0 else None

        # Alinhamento e borda por célula
        alignments = [center, center, left, center, center, center, right, center, center]
        for col_idx, align in enumerate(alignments, start=1):
            cell = ws.cell(row=excel_row, column=col_idx)
            cell.alignment = align
            cell.border    = cell_border
            if fill:
                cell.fill = fill

        ws.row_dimensions[excel_row].height = 18
        row_num += 1

    # Congelar linha do cabeçalho
    ws.freeze_panes = "A2"

    # Salvar em memória
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return Response(
        content=output.read(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=resultado_lote.xlsx"},
    )


class SafeStaticFiles(StaticFiles):
    async def get_response(self, path: str, scope):
        allowed_root_files = {
            "",
            "index.html",
            "dashboard.html",
            "history.html",
            "notes.html",
            "style.css",
            "script.js"
        }
        allowed_directories = ("css/", "js/", "instrucoes_api/")
        
        norm_path = path.replace("\\", "/").strip("/")
        
        is_allowed = (
            norm_path in allowed_root_files or
            any(norm_path.startswith(d_pref) for d_pref in allowed_directories)
        )
        
        if not is_allowed:
            from starlette.exceptions import HTTPException as StarletteHTTPException
            raise StarletteHTTPException(status_code=404, detail="Not Found")
            
        return await super().get_response(path, scope)

# Mount static files last
app.mount("/", SafeStaticFiles(directory=".", html=True), name="static")

if __name__ == "__main__":
    import os
    def get_free_port():
        s = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
        s.bind(('', 0))
        port = s.getsockname()[1]
        s.close()
        return port

    port = int(os.environ.get("PORT", get_free_port()))
    # No emoji in print to avoid Windows console errors
    print(f"\n\nSERVER RUNNING ON: http://localhost:{port}\n\n")
    uvicorn.run(app, host="0.0.0.0", port=port)
